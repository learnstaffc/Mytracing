"""生成 README 演示用的虚构示例数据（examples/sample_data.xlsx）。

数据全部是编造的整齐数字，不对应任何真实账户，仅用于展示图表效果。
ZSBank / WeFinace / A-inv 单位为人民币，US-inv 单位为美元，Credit 为负债。

用法：
    python examples/make_sample_data.py
    python main.py --file examples/sample_data.xlsx --sheet record
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

HEADERS = ["日期", "ZSBank", "WeFinace", "A-inv", "US-inv", "Credit"]
FX_US_INV = 6.6  # US-inv 折算汇率，与 src/config.py 保持一致

# (日期, ZSBank, WeFinace, A-inv, US-inv(USD), Credit) —— 每月 1 号一条，纯虚构
SAMPLE_ROWS = [
    ("2025-01-01",  60000, 35000, 40000, 4000,    0),
    ("2025-02-01",  63000, 35200, 41500, 4100,    0),
    ("2025-03-01",  66000, 35500, 40800, 4050,  800),
    ("2025-04-01",  69000, 36000, 39000, 4150,    0),
    ("2025-05-01",  71500, 36500, 37800, 4300,    0),
    ("2025-06-01",  74000, 37000, 39200, 4400, 1200),
    ("2025-07-01",  76500, 37500, 40500, 4450,    0),
    ("2025-08-01",  80000, 38000, 38800, 4600,    0),
    ("2025-09-01",  83000, 38500, 37600, 4700, 2000),
    ("2025-10-01",  86000, 39000, 39500, 4800,    0),
    ("2025-11-01",  90000, 39500, 41200, 4950,    0),
    ("2025-12-01",  95000, 40000, 43000, 5100, 3000),
    ("2026-01-01", 100000, 40500, 41800, 5250,    0),
    ("2026-02-01", 106000, 41000, 44200, 5400, 1500),
    ("2026-03-01", 112000, 41500, 45600, 5600,    0),
    ("2026-04-01", 118000, 42000, 47000, 5750,    0),
    ("2026-05-01", 124000, 42500, 45800, 5900, 2500),
    ("2026-06-01", 131000, 43000, 48200, 6050,    0),
]


def main() -> None:
    target = Path(__file__).resolve().parent / "sample_data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "record"

    for col, name in enumerate(HEADERS + ["Sum"], start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    for r, (d, zs, wf, ai, us, cr) in enumerate(SAMPLE_ROWS, start=2):
        ws.cell(row=r, column=1, value=d).number_format = "yyyy-mm-dd"
        for col, value in enumerate((zs, wf, ai), start=2):
            ws.cell(row=r, column=col, value=value).number_format = "#,##0.00"
        ws.cell(row=r, column=5, value=us).number_format = "#,##0"
        ws.cell(row=r, column=6, value=cr).number_format = "#,##0"
        # Sum = ZSBank + WeFinace + A-inv + US-inv*汇率 - Credit
        ws.cell(row=r, column=7,
                value=f"=B{r}+C{r}+D{r}+E{r}*{FX_US_INV}-F{r}").number_format = "#,##0.00"

    for letter, w in {"A": 12, "B": 12, "C": 12, "D": 12, "E": 10, "F": 10, "G": 14}.items():
        ws.column_dimensions[letter].width = w

    wb.save(target)
    print(f"示例数据已生成: {target}")


if __name__ == "__main__":
    main()
