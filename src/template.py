"""生成示例资产负债模板（写入用户指定路径，供照着填写）。

列与你的实际文件一致：日期 / ZSBank / WeFinace / A-inv / US-inv / Credit / Sum
其中 Sum 列按公式自动生成：Sum = ZSBank + WeFinace + A-inv + US-inv*6.6 - Credit
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

HEADERS = ["日期", "ZSBank", "WeFinace", "A-inv", "US-inv", "Credit", "Sum"]

# 示例数据（每月 1 号记一次，纯虚构的整齐数字，仅供演示格式）：
# ZSBank/WeFinace/A-inv 为人民币，US-inv 为美元，Credit 为负债
SAMPLE_ROWS = [
    ("2026-01-01", 60000, 35000, 40000, 4000, 0),
    ("2026-02-01", 62000, 36000, 41000, 4200, 2000),
    ("2026-03-01", 65000, 37000, 42000, 4500, 3000),
]


def create_template(path: str) -> str:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = config.SHEET

    # 表头
    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    # 示例数据行，Sum 列用公式
    for r, (d, zs, wf, ai, us, cr) in enumerate(SAMPLE_ROWS, start=2):
        ws.cell(row=r, column=1, value=d).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2, value=zs).number_format = "#,##0.00"
        ws.cell(row=r, column=3, value=wf).number_format = "#,##0.00"
        ws.cell(row=r, column=4, value=ai).number_format = "#,##0.00"
        ws.cell(row=r, column=5, value=us).number_format = "#,##0"
        ws.cell(row=r, column=6, value=cr).number_format = "#,##0"
        # Sum = ZSBank + WeFinace + A-inv + US-inv*6.6 - Credit
        ws.cell(row=r, column=7, value=f"=B{r}+C{r}+D{r}+E{r}*{config.FX_RATES.get('US-inv', 6.6)}-F{r}"
                ).number_format = "#,##0.00"

    # 列宽
    widths = {"A": 12, "B": 12, "C": 12, "D": 12, "E": 10, "F": 10, "G": 14}
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w

    wb.save(target)
    return str(target)
